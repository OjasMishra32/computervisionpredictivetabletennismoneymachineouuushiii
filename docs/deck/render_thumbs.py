#!/usr/bin/env python3
"""Slide thumbnails for docs/deck/courtside.pptx without LibreOffice or Keynote.

    .venv/bin/python docs/deck/render_thumbs.py      # docs/deck/thumbs/slide_NN.png + contact_sheet.png

A small PIL renderer for exactly what docs/deck/build_deck.py draws: solid rectangles / rounded rectangles / ovals,
straight connectors, text boxes (word wrap, alignment, anchors, spacing), hairline tables, pictures and movie poster
frames. It is a preview, not PowerPoint: kerning and line breaks can differ by a word. It also reports any text that
does not fit its box (exit code 1), which is the layout check the deck build cannot do on its own.
"""
from __future__ import annotations

import io
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.dml import MSO_FILL
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.shapes.connector import Connector

DECK = Path(__file__).resolve().parent
ROOT = DECK.parents[1]
PPTX = DECK / "courtside.pptx"
OUT = DECK / "thumbs"
FONTS = ROOT / "results/viz/v60_assets/fonts"
PX = 120                      # pixels per inch: 1600 x 900
EMU = 914400

_cache: dict = {}


def font(name: str, bold: bool, size_px: float):
    key = (name, bold, round(size_px))
    if key in _cache:
        return _cache[key]
    if name and name.lower().startswith(("menlo", "courier", "consolas")):
        f = ImageFont.truetype("/System/Library/Fonts/Menlo.ttc", max(1, round(size_px)), index=1 if bold else 0)
    else:
        f = ImageFont.truetype(str(FONTS / ("Inter_600SemiBold.ttf" if bold else "Inter_400Regular.ttf")), max(1, round(size_px)))
    _cache[key] = f
    return f


def px(emu) -> float:
    return emu / EMU * PX


def color_of(fc, default="000000"):
    try:
        return "#" + str(fc.rgb)
    except Exception:  # noqa: BLE001
        return "#" + default


def run_style(run, para_default_size=18):
    f = run.font
    size = f.size.pt if f.size else para_default_size
    return {"size": size, "bold": bool(f.bold), "name": f.name or "Inter", "color": color_of(f.color, "000000")}


def layout_text(tf, w_px):
    """Lines of (segments, height, align) for a text frame of inner width w_px; segments = (x, text, style)."""
    lines = []
    for para in tf.paragraphs:
        runs = [(r.text, run_style(r)) for r in para.runs if r.text]
        ls = para.line_spacing if isinstance(para.line_spacing, float) else 1.0
        before = para.space_before.pt / 72 * PX if para.space_before is not None else 0
        after = para.space_after.pt / 72 * PX if para.space_after is not None else 0
        align = para.alignment or PP_ALIGN.LEFT
        if not runs:
            lines.append({"segs": [], "h": 18 / 72 * PX * 1.2, "align": align, "before": before, "after": after, "w": 0})
            continue
        tokens = []
        for t, st in runs:
            parts = t.replace("\n", " ").split(" ")
            for i, p in enumerate(parts):
                tokens.append((p, st, i < len(parts) - 1))
        cur, cur_w, cur_h, first = [], 0.0, 0.0, True
        for word, st, space_after in tokens:
            f = font(st["name"], st["bold"], st["size"] / 72 * PX)
            ww = f.getlength(word)
            sw = f.getlength(" ") if space_after else 0
            if cur and cur_w + ww > w_px + 0.5 and word:
                lines.append({"segs": cur, "h": cur_h, "align": align, "before": before if first else 0, "after": 0,
                              "w": cur_w})
                first = False
                cur, cur_w, cur_h = [], 0.0, 0.0
                cur_w_trim = 0
            cur.append((cur_w, word + (" " if space_after else ""), st))
            cur_w += ww + sw
            cur_h = max(cur_h, st["size"] / 72 * PX * 1.2 * ls)
        lines.append({"segs": cur, "h": cur_h, "align": align, "before": before if first else 0, "after": after,
                      "w": cur_w})
    return lines


def draw_text(d, tf, x, y, w, h, anchor_default=MSO_ANCHOR.TOP, insets=None):
    bp = tf._txBody.find(qn("a:bodyPr"))
    li = int(bp.get("lIns", 91440)) if bp is not None else 91440
    ti = int(bp.get("tIns", 45720)) if bp is not None else 45720
    ri = int(bp.get("rIns", 91440)) if bp is not None else 91440
    bi = int(bp.get("bIns", 45720)) if bp is not None else 45720
    if insets:
        li, ti, ri, bi = insets
    ix, iy, iw, ih = x + px(li), y + px(ti), w - px(li) - px(ri), h - px(ti) - px(bi)
    wrap = tf.word_wrap is not False
    lines = layout_text(tf, iw if wrap else 1e9)
    total = sum(L["h"] + L["before"] + L["after"] for L in lines)
    anchor = tf.vertical_anchor or anchor_default
    cy = iy if anchor == MSO_ANCHOR.TOP else (iy + (ih - total) / 2 if anchor == MSO_ANCHOR.MIDDLE else iy + ih - total)
    for L in lines:
        cy += L["before"]
        lw = L["w"] - (font(L["segs"][-1][2]["name"], L["segs"][-1][2]["bold"], L["segs"][-1][2]["size"] / 72 * PX).getlength(" ")
                       if L["segs"] and L["segs"][-1][1].endswith(" ") else 0)
        ox = ix if L["align"] in (PP_ALIGN.LEFT, None) else (ix + (iw - lw) / 2 if L["align"] == PP_ALIGN.CENTER else ix + iw - lw)
        for sx, t, st in L["segs"]:
            f = font(st["name"], st["bold"], st["size"] / 72 * PX)
            asc = f.getmetrics()[0]
            d.text((ox + sx, cy + (L["h"] - st["size"] / 72 * PX * 1.2) + st["size"] / 72 * PX * 0.95 - asc + 2), t,
                   font=f, fill=st["color"])
        cy += L["h"] + L["after"]
    widest = max((L["w"] for L in lines), default=0)
    over_h = total - ih
    over_w = widest - iw if not wrap else max((L["w"] - iw for L in lines if len(L["segs"]) == 1), default=0)
    return over_h, over_w


def fill_hex(shape):
    try:
        if shape.fill.type == MSO_FILL.SOLID:
            return "#" + str(shape.fill.fore_color.rgb)
    except Exception:  # noqa: BLE001
        pass
    return None


def line_hex(shape):
    try:
        if shape.line.fill.type == MSO_FILL.SOLID:
            return "#" + str(shape.line.color.rgb), max(1, round(shape.line.width.pt / 72 * PX))
    except Exception:  # noqa: BLE001
        pass
    return None, 0


def render(prs, k, slide):
    W, H = round(px(prs.slide_width)), round(px(prs.slide_height))
    bg = "#FFFFFF"
    try:
        if slide.background.fill.type == MSO_FILL.SOLID:
            bg = "#" + str(slide.background.fill.fore_color.rgb)
    except Exception:  # noqa: BLE001
        pass
    im = Image.new("RGB", (W, H), bg)
    d = ImageDraw.Draw(im)
    issues = []
    for sh in slide.shapes:
        x, y, w, h = px(sh.left or 0), px(sh.top or 0), px(sh.width or 0), px(sh.height or 0)
        if isinstance(sh, Connector):
            col, lw = line_hex(sh)
            d.line([(px(sh.begin_x), px(sh.begin_y)), (px(sh.end_x), px(sh.end_y))], fill=col or "#888888", width=lw or 1)
            continue
        if sh.shape_type in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.MEDIA):
            blip = sh._element.find(".//" + qn("a:blip"))
            rid = blip.get(qn("r:embed")) if blip is not None else None
            if rid:
                pic = Image.open(io.BytesIO(slide.part.related_part(rid).blob)).convert("RGB")
                im.paste(pic.resize((max(1, round(w)), max(1, round(h))), Image.LANCZOS), (round(x), round(y)))
                if sh.shape_type == MSO_SHAPE_TYPE.MEDIA:   # a play badge marks the movie
                    cx, cy, r = x + w / 2, y + h / 2, 0.28 * PX
                    d.ellipse([cx - r, cy - r, cx + r, cy + r], fill="#000000")
                    d.polygon([(cx - r / 3, cy - r / 2), (cx - r / 3, cy + r / 2), (cx + r / 2, cy)], fill="#FFFFFF")
            continue
        if sh.has_table if hasattr(sh, "has_table") else False:
            tbl = sh.table
            cy = y
            for i, row in enumerate(tbl.rows):
                rh = px(row.height)
                cx = x
                for j, cell in enumerate(row.cells):
                    cw = px(tbl.columns[j].width)
                    ins = (cell.margin_left, cell.margin_top, cell.margin_right, cell.margin_bottom)
                    oh, ow = draw_text(d, cell.text_frame, cx, cy, cw, rh, cell.vertical_anchor or MSO_ANCHOR.TOP, ins)
                    if oh > 3 or ow > 3:
                        issues.append(f"slide {k} table '{sh.name}' cell ({i},{j}) overflows by {max(oh, ow) / PX:.2f} in: "
                                      f"{cell.text_frame.text[:40]!r}")
                    ln = cell._tc.tcPr.find(qn("a:lnB")) if cell._tc.tcPr is not None else None
                    clr = ln.find(".//" + qn("a:srgbClr")) if ln is not None else None
                    if clr is not None:
                        d.line([(cx, cy + rh), (cx + cw, cy + rh)], fill="#" + clr.get("val"), width=1)
                    cx += cw
                cy += rh
            continue
        if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
            fcol = fill_hex(sh)
            lcol, lw = line_hex(sh)
            kind = sh.auto_shape_type
            box = [x, y, x + w, y + h]
            if kind == MSO_SHAPE.OVAL:
                d.ellipse(box, fill=fcol, outline=lcol, width=lw)
            elif kind == MSO_SHAPE.ROUNDED_RECTANGLE:
                r = sh.adjustments[0] * min(w, h) if len(sh.adjustments) else 0.1 * min(w, h)
                d.rounded_rectangle(box, radius=r, fill=fcol, outline=lcol, width=lw)
            else:
                d.rectangle(box, fill=fcol, outline=lcol, width=lw)
        if sh.has_text_frame and sh.text_frame.text.strip():
            oh, ow = draw_text(d, sh.text_frame, x, y, w, h)
            if oh > 3 or ow > 3:
                issues.append(f"slide {k} '{sh.name}' overflows by {max(oh, ow) / PX:.2f} in: {sh.text_frame.text[:50]!r}")
        if x + w > W + 2 or y + h > H + 2 or x < -2 or y < -2:
            issues.append(f"slide {k} '{sh.name}' is off the slide")
    return im, issues


def main() -> int:
    prs = Presentation(str(PPTX))
    OUT.mkdir(exist_ok=True)
    for old in OUT.glob("slide_*.png"):
        old.unlink()
    thumbs, issues = [], []
    for k, slide in enumerate(prs.slides, 1):
        im, iss = render(prs, k, slide)
        issues += iss
        small = im.resize((960, 540), Image.LANCZOS)
        small.save(OUT / f"slide_{k:02d}.png", optimize=True)
        thumbs.append(small)
    cols, tw, th, pad = 4, 480, 270, 12
    rows = (len(thumbs) + cols - 1) // cols
    sheet = Image.new("RGB", (cols * (tw + pad) + pad, rows * (th + pad) + pad), "#8E8E93")
    for i, t in enumerate(thumbs):
        sheet.paste(t.resize((tw, th), Image.LANCZOS), (pad + (i % cols) * (tw + pad), pad + (i // cols) * (th + pad)))
    sheet.save(OUT / "contact_sheet.png", optimize=True)
    print(f"wrote {len(thumbs)} thumbnails + contact_sheet.png to {OUT.relative_to(ROOT)}")
    for x in issues:
        print("LAYOUT:", x)
    print("layout:", "PASS" if not issues else f"{len(issues)} issue(s)")
    return 1 if issues else 0


if __name__ == "__main__":
    sys.exit(main())
